"""服务端 Excel 导出:复刻 5 个 IE 分析类模板前端组件的 exportExcel 逻辑。

前端编辑器里的「导出 Excel」由各模板 Vue 组件用 exceljs 在浏览器生成;
知识库问答等外部调用方拿不到浏览器环境,只有图表记录里的 chartData。
本模块用 openpyxl(+Pillow 画连线叠加图)按同样的布局/合并/配色规则
服务端重建同款工作簿,保证两个入口导出的文件一致。

支持模板: man-machine / two-hand / joint-operation / process-flow / therblig
入口: build_chart_excel(template_id, chart_data) -> (xlsx_bytes, filename)
"""

from __future__ import annotations

import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor, TwoCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image as PILImage
from PIL import ImageDraw

# ── 通用样式(与前端 exceljs 版一致) ──────────────────────────────
_THIN = Side(style="thin", color="FF777777")
ALL_BORDER = Border(top=_THIN, left=_THIN, right=_THIN, bottom=_THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
HEADER_FILL = PatternFill(fill_type="solid", fgColor="FFE8E8E8")
SUB_HEADER_FILL = PatternFill(fill_type="solid", fgColor="FFF0F0F0")
LABEL_FILL_SLATE = PatternFill(fill_type="solid", fgColor="FFF1F5F9")
LABEL_FILL_GRAY = PatternFill(fill_type="solid", fgColor="FFF5F5F5")
WORK_FILL = PatternFill(fill_type="solid", fgColor="FFA0A0A0")
IDLE_FILL = PatternFill(fill_type="solid", fgColor="FFFFFFFF")

GREEN = "FF16A34A"
RED = "FFDC2626"
GRAY = "FF94A3B8"


def _ws(title: str):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    return wb, ws


def _set_widths(ws, widths: list[float]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _cell(ws, r: int, c: int, value=None, *, font: Font | None = None,
          fill: PatternFill | None = None, align: Alignment | None = CENTER):
    cell = ws.cell(row=r, column=c)
    if value is not None:
        cell.value = value
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    if align is not None:
        cell.alignment = align
    return cell


def _merge(ws, r1: int, c1: int, r2: int, c2: int) -> None:
    # 列数很少时 split_cols 可能给出退化区间(e<s),按单格处理不合并
    if r2 < r1 or c2 < c1 or (r1 == r2 and c1 == c2):
        return
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def _border_all(ws, last_row: int, last_col: int) -> None:
    for r in range(1, last_row + 1):
        for c in range(1, last_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = ALL_BORDER
            if cell.alignment is None or cell.alignment.horizontal is None:
                cell.alignment = CENTER


# ── 连线叠加图(前端用 canvas,这里用 Pillow;透明底 PNG) ──────────
def _col_px(width: float) -> int:
    # 与前端一致:Excel 字符宽 → 像素的经验换算
    return round(width * 7 + 5)


def _row_px(height: float) -> int:
    return round(height * 96 / 72)


def _polyline_png(col_widths_px: list[int], row_heights_px: list[int],
                  col_indexes: list[int], *, color: tuple[int, int, int],
                  line_width: float, dot_radius: float | None = None) -> bytes | None:
    """按「每行落在第 col_indexes[i] 列的中心」画折线,返回透明 PNG。"""
    if len(col_indexes) < 2:
        return None
    total_w = sum(col_widths_px)
    total_h = sum(row_heights_px)
    scale = 2
    img = PILImage.new("RGBA", (total_w * scale, total_h * scale), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)

    col_start = [0]
    for w in col_widths_px[:-1]:
        col_start.append(col_start[-1] + w)
    row_start = [0]
    for h in row_heights_px[:-1]:
        row_start.append(row_start[-1] + h)

    pts = []
    for i, ci in enumerate(col_indexes):
        ci = max(0, min(ci, len(col_widths_px) - 1))
        x = (col_start[ci] + col_widths_px[ci] / 2) * scale
        y = (row_start[i] + row_heights_px[i] / 2) * scale
        pts.append((x, y))

    rgba = (*color, 255)
    draw.line(pts, fill=rgba, width=max(1, round(line_width * scale)), joint="curve")
    if dot_radius:
        r = dot_radius * scale
        for x, y in pts:
            draw.ellipse((x - r, y - r, x + r, y + r), fill=rgba)

    buf = BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def _overlay_image(ws, png: bytes, *, tl_col: int, tl_row: int, br_col: int, br_row: int) -> None:
    """等价 exceljs 的 addImage(tl/br, editAs='oneCell'),行列均为 0 基。"""
    img = XLImage(BytesIO(png))
    anchor = TwoCellAnchor(editAs="oneCell")
    anchor._from = AnchorMarker(col=tl_col, colOff=0, row=tl_row, rowOff=0)
    anchor.to = AnchorMarker(col=br_col, colOff=0, row=br_row, rowOff=0)
    img.anchor = anchor
    ws.add_image(img)


def _num(v, default=0):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return default
    return int(f) if f.is_integer() else f


def _save(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ═══════════════════════ 人机作业分析 man-machine ═══════════════════════
def _build_man_machine(data: dict) -> tuple[bytes, str]:
    header = data.get("headerInfo") or {}
    machines = [m for m in (data.get("machines") or []) if isinstance(m, dict)]
    rows = [r for r in (data.get("rows") or []) if isinstance(r, dict)]

    def get_cell(row: dict, mid: str) -> dict:
        c = (row.get("cells") or {}).get(mid)
        return c if isinstance(c, dict) else {"desc": "", "state": "idle"}

    wb, ws = _ws("人机作业分析")
    m_count = len(machines)
    total_cols = 3 + m_count * 2
    _set_widths(ws, [10, 22, 8] + [22, 8] * m_count)

    bold = Font(bold=True)

    # 标题
    _merge(ws, 1, 1, 1, total_cols)
    _cell(ws, 1, 1, "人机作业分析图", font=Font(bold=True, size=14))
    ws.row_dimensions[1].height = 30

    # 信息区 Row2-4
    _cell(ws, 2, 1, "作业名称", font=bold, fill=LABEL_FILL_SLATE)
    _cell(ws, 2, 2, header.get("taskName") or "")
    if total_cols > 2:
        _merge(ws, 2, 2, 2, total_cols)
    _cell(ws, 3, 1, "操作者", font=bold, fill=LABEL_FILL_SLATE)
    _cell(ws, 3, 2, header.get("worker") or "")
    _cell(ws, 3, 3, "机器", font=bold, fill=LABEL_FILL_SLATE)
    _cell(ws, 3, 4, "、".join(str(m.get("name") or "") for m in machines))
    if total_cols > 4:
        _merge(ws, 3, 4, 3, total_cols)
    _cell(ws, 4, 1, "日期", font=bold, fill=LABEL_FILL_SLATE)
    _cell(ws, 4, 2, header.get("date") or "")
    _cell(ws, 4, 3, "研究者", font=bold, fill=LABEL_FILL_SLATE)
    _cell(ws, 4, 4, header.get("analyst") or "")
    if total_cols > 4:
        _merge(ws, 4, 4, 4, total_cols)
    for r in range(2, 5):
        for c in range(1, total_cols + 1):
            ws.cell(row=r, column=c).alignment = CENTER

    # 表头 Row5-6
    _cell(ws, 5, 1, "时间")
    _merge(ws, 5, 1, 6, 1)
    _cell(ws, 5, 2, header.get("worker") or "操作者")
    _merge(ws, 5, 2, 5, 3)
    _cell(ws, 6, 2, "描述")
    _cell(ws, 6, 3, "状态")
    for mi, m in enumerate(machines):
        base = 4 + mi * 2
        _cell(ws, 5, base, m.get("name") or "")
        _merge(ws, 5, base, 5, base + 1)
        _cell(ws, 6, base, "描述")
        _cell(ws, 6, base + 1, "状态")
    for r in (5, 6):
        for c in range(1, total_cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = bold
            cell.alignment = CENTER
            cell.fill = HEADER_FILL

    # 数据行
    start_row = 7
    cum = 0
    for idx, row in enumerate(rows):
        r = start_row + idx
        cum += _num(row.get("time"))
        _cell(ws, r, 1, cum)
        _cell(ws, r, 2, row.get("workerDesc") or "")
        ws.cell(row=r, column=3).fill = WORK_FILL if row.get("workerState") == "work" else IDLE_FILL
        for mi, m in enumerate(machines):
            base = 4 + mi * 2
            cell = get_cell(row, m.get("id"))
            _cell(ws, r, base, cell.get("desc") or "")
            ws.cell(row=r, column=base + 1).fill = (
                WORK_FILL if cell.get("state") == "work" else IDLE_FILL
            )
        ws.row_dimensions[r].height = 22
        for c in range(1, total_cols + 1):
            ws.cell(row=r, column=c).alignment = CENTER

    last_data_row = start_row + max(len(rows), 1) - 1

    # 合并:人描述列 / 机器描述列(desc+state 连续相同)
    def merge_groups(key_fn, col: int) -> None:
        i = 0
        while i < len(rows):
            span = 1
            while i + span < len(rows) and key_fn(rows[i + span]) == key_fn(rows[i]):
                span += 1
            if span > 1:
                _merge(ws, start_row + i, col, start_row + i + span - 1, col)
            i += span

    merge_groups(lambda r: (r.get("workerDesc"), r.get("workerState")), 2)
    for mi, m in enumerate(machines):
        mid = m.get("id")
        merge_groups(lambda r, _mid=mid: (get_cell(r, _mid).get("desc"),
                                          get_cell(r, _mid).get("state")), 4 + mi * 2)
    # 合并:状态列(连续相同状态)
    merge_groups(lambda r: r.get("workerState"), 3)
    for mi, m in enumerate(machines):
        mid = m.get("id")
        merge_groups(lambda r, _mid=mid: get_cell(r, _mid).get("state"), 5 + mi * 2)

    # 统计
    cycle = sum(_num(r.get("time")) for r in rows)
    worker_work = sum(_num(r.get("time")) for r in rows if r.get("workerState") == "work")

    def eff(work) -> float:
        return round(work / cycle * 100, 1) if cycle > 0 else 0

    sr = last_data_row + 1
    _merge(ws, sr, 1, sr, total_cols)
    _cell(ws, sr, 1, "统计汇总", font=Font(bold=True, size=12))

    hr = sr + 1
    for i, h in enumerate(["项目", "周程/s", "工作时间/s", "空闲时间/s", "利用率/%"]):
        _cell(ws, hr, i + 1, h, font=bold, fill=HEADER_FILL)
    if total_cols > 5:
        _merge(ws, hr, 5, hr, total_cols)

    def write_stat(r: int, label: str, work, idle, e) -> None:
        _cell(ws, r, 1, label, font=bold)
        _cell(ws, r, 2, cycle)
        _cell(ws, r, 3, work)
        _cell(ws, r, 4, idle)
        _cell(ws, r, 5, f"{e}%",
              font=Font(bold=True, color=GREEN if e >= 80 else RED))
        for c in range(1, max(5, total_cols) + 1):
            ws.cell(row=r, column=c).alignment = CENTER
        if total_cols > 5:
            _merge(ws, r, 5, r, total_cols)

    write_stat(hr + 1, header.get("worker") or "操作者",
               worker_work, cycle - worker_work, eff(worker_work))
    for mi, m in enumerate(machines):
        mid = m.get("id")
        m_work = sum(_num(r.get("time")) for r in rows if get_cell(r, mid).get("state") == "work")
        write_stat(hr + 2 + mi, m.get("name") or "", m_work, cycle - m_work, eff(m_work))

    _border_all(ws, hr + 1 + len(machines), total_cols)
    name = header.get("taskName") or "导出"
    return _save(wb), f"人机作业分析_{name}.xlsx"


# ═══════════════════════ 双手作业分析 two-hand ═══════════════════════
_TH_SYMBOL_KEYS = ["operation", "transport", "delay", "hold"]
_TH_SYMBOL_DISPLAY = {"operation": "○", "transport": "⇨", "delay": "D", "hold": "▽"}
_TH_OLD_SYMBOL = {"operation": "operation", "transport": "transport", "inspect": "hold",
                  "wait": "delay", "hold": "hold", "delay": "delay"}


def _build_two_hand(data: dict) -> tuple[bytes, str]:
    header = data.get("headerInfo") or {}
    steps = []
    for s in data.get("steps") or []:
        if not isinstance(s, dict):
            continue
        steps.append({
            "leftDesc": s.get("leftDesc") or "",
            "rightDesc": s.get("rightDesc") or "",
            "leftSymbol": _TH_OLD_SYMBOL.get(s.get("leftSymbol"), "operation"),
            "rightSymbol": _TH_OLD_SYMBOL.get(s.get("rightSymbol"), "operation"),
        })

    wb, ws = _ws("双手作业分析图")
    widths = [22, 7, 7, 7, 7, 7, 7, 7, 7, 22]
    _set_widths(ws, widths)

    cjk = Font(name="宋体", size=11)
    cjk_bold = Font(name="宋体", size=11, bold=True)
    sym_font = Font(name="宋体", size=11, bold=True)

    _merge(ws, 1, 1, 1, 10)
    _cell(ws, 1, 1, "双手作业分析图", font=Font(name="宋体", size=16, bold=True))
    ws.row_dimensions[1].height = 34

    _merge(ws, 2, 1, 3, 1)
    _cell(ws, 2, 1, "左手", font=cjk_bold, fill=HEADER_FILL)
    _merge(ws, 2, 2, 2, 5)
    _cell(ws, 2, 2, "符号", font=cjk_bold, fill=HEADER_FILL)
    _merge(ws, 2, 6, 2, 9)
    _cell(ws, 2, 6, "符号", font=cjk_bold, fill=HEADER_FILL)
    _merge(ws, 2, 10, 3, 10)
    _cell(ws, 2, 10, "右手", font=cjk_bold, fill=HEADER_FILL)
    for i, h in enumerate(["○", "⇨", "D", "▽", "○", "⇨", "D", "▽"]):
        _cell(ws, 3, i + 2, h, font=sym_font, fill=HEADER_FILL)
    for r in (2, 3):
        ws.row_dimensions[r].height = 24
        for c in range(1, 11):
            cell = ws.cell(row=r, column=c)
            cell.font = sym_font if (r == 3 and 2 <= c <= 9) else cjk_bold
            cell.fill = HEADER_FILL
            cell.alignment = CENTER

    data_start = 4
    row_h = 26
    for idx, step in enumerate(steps):
        r = data_start + idx
        _cell(ws, r, 1, step["leftDesc"] or "—")
        _cell(ws, r, 10, step["rightDesc"] or "—")
        ws.row_dimensions[r].height = row_h
        for c in range(1, 11):
            cell = ws.cell(row=r, column=c)
            cell.font = cjk
            cell.alignment = CENTER

    # 左右两侧连线 + 圆点叠加图(B-E / F-I)
    if len(steps) >= 2:
        sym_widths_px = [_col_px(widths[c - 1]) for c in (2, 3, 4, 5)]
        row_heights_px = [_row_px(row_h)] * len(steps)
        for side, tl_col, br_col in (("leftSymbol", 1, 5), ("rightSymbol", 5, 9)):
            cols = [_TH_SYMBOL_KEYS.index(s[side]) for s in steps]
            png = _polyline_png(sym_widths_px, row_heights_px, cols,
                                color=(51, 51, 51), line_width=1.5, dot_radius=3.5)
            if png:
                _overlay_image(ws, png, tl_col=tl_col, tl_row=data_start - 1,
                               br_col=br_col, br_row=data_start - 1 + len(steps))

    last_data_row = data_start + max(len(steps), 1) - 1

    # 统计表
    stats_start = last_data_row + 1
    for (c1, c2), label in zip([(1, 2), (3, 4), (5, 7), (8, 10)],
                               ["统计内容", "符号", "左手", "右手"]):
        _merge(ws, stats_start, c1, stats_start, c2)
        _cell(ws, stats_start, c1, label, font=cjk_bold, fill=HEADER_FILL)
    ws.row_dimensions[stats_start].height = 24
    for c in range(1, 11):
        cell = ws.cell(row=stats_start, column=c)
        cell.font = cjk_bold
        cell.fill = HEADER_FILL
        cell.alignment = CENTER

    left_counts = {k: 0 for k in _TH_SYMBOL_KEYS}
    right_counts = {k: 0 for k in _TH_SYMBOL_KEYS}
    for s in steps:
        left_counts[s["leftSymbol"]] += 1
        right_counts[s["rightSymbol"]] += 1

    sym_row = stats_start + 1
    _merge(ws, sym_row, 1, sym_row + 3, 2)
    _cell(ws, sym_row, 1, "次数", font=cjk_bold)
    for idx, sym in enumerate(_TH_SYMBOL_KEYS):
        r = sym_row + idx
        _merge(ws, r, 3, r, 4)
        _cell(ws, r, 3, _TH_SYMBOL_DISPLAY[sym], font=sym_font)
        _merge(ws, r, 5, r, 7)
        _cell(ws, r, 5, left_counts[sym], font=cjk)
        _merge(ws, r, 8, r, 10)
        _cell(ws, r, 8, right_counts[sym], font=cjk)
        ws.row_dimensions[r].height = 22

    total_row = sym_row + 4
    _merge(ws, total_row, 1, total_row, 2)
    _cell(ws, total_row, 1, "合计", font=cjk_bold)
    _merge(ws, total_row, 3, total_row, 4)
    _merge(ws, total_row, 5, total_row, 7)
    _cell(ws, total_row, 5, sum(left_counts.values()), font=cjk_bold)
    _merge(ws, total_row, 8, total_row, 10)
    _cell(ws, total_row, 8, sum(right_counts.values()), font=cjk_bold)
    ws.row_dimensions[total_row].height = 24

    _border_all(ws, total_row, 10)
    name = header.get("stationName") or "导出"
    return _save(wb), f"双手作业分析_{name}.xlsx"


# ═══════════════════════ 联合作业分析 joint-operation ═══════════════════════
def _build_joint_operation(data: dict) -> tuple[bytes, str]:
    header = data.get("headerInfo") or {}
    columns = [c for c in (data.get("columns") or []) if isinstance(c, dict)]
    time_rows = [r for r in (data.get("timeRows") or []) if isinstance(r, dict)]

    def get_cell(row: dict, cid: str) -> dict:
        c = (row.get("cells") or {}).get(cid)
        return c if isinstance(c, dict) else {"desc": "", "duration": 0, "isWork": False}

    wb, ws = _ws("联合作业分析")
    n = len(columns)
    total_cols = 1 + n * 2
    _set_widths(ws, [10] + [18, 5] * n)

    bold = Font(bold=True)
    row_i = 1

    _merge(ws, row_i, 1, row_i, total_cols)
    _cell(ws, row_i, 1, "联合作业分析图", font=Font(bold=True, size=14))
    ws.row_dimensions[row_i].height = 30
    row_i += 1

    def split_cols(parts: int) -> list[tuple[int, int]]:
        segs = []
        for i in range(parts):
            s = 1 + (i * total_cols) // parts
            e = ((i + 1) * total_cols) // parts
            segs.append((s, e))
        return segs

    # 信息行 1:作业名称 | 操作者
    h1 = split_cols(4)
    for (s, e), value, is_label in [
        (h1[0], "作业名称", True), (h1[1], header.get("taskName") or "—", False),
        (h1[2], "操作者", True), (h1[3], header.get("operator") or "—", False),
    ]:
        _merge(ws, row_i, s, row_i, e)
        _cell(ws, row_i, s, value, font=bold if is_label else None,
              fill=LABEL_FILL_GRAY if is_label else None)
    ws.row_dimensions[row_i].height = 24
    row_i += 1

    # 信息行 2:日期 | 地点 | 备注
    h2 = split_cols(6)
    kv = [("日期", header.get("date") or "—"), ("地点", header.get("location") or "—"),
          ("备注", header.get("remarks") or "—")]
    for i, (k, v) in enumerate(kv):
        ks, ke = h2[i * 2]
        vs, ve = h2[i * 2 + 1]
        _merge(ws, row_i, ks, row_i, ke)
        _cell(ws, row_i, ks, k, font=bold, fill=LABEL_FILL_GRAY)
        _merge(ws, row_i, vs, row_i, ve)
        _cell(ws, row_i, vs, v)
    ws.row_dimensions[row_i].height = 24
    row_i += 1

    # 数据表头两行
    _cell(ws, row_i, 1, "时间/s", font=bold, fill=HEADER_FILL)
    _merge(ws, row_i, 1, row_i + 1, 1)
    for idx, col in enumerate(columns):
        sc = 2 + idx * 2
        _merge(ws, row_i, sc, row_i, sc + 1)
        _cell(ws, row_i, sc, col.get("name") or "", font=bold, fill=HEADER_FILL)
    ws.row_dimensions[row_i].height = 26
    row_i += 1
    for idx in range(n):
        sc = 2 + idx * 2
        _cell(ws, row_i, sc, "描述", font=Font(bold=True, size=10), fill=SUB_HEADER_FILL)
        _cell(ws, row_i, sc + 1, "时间", font=Font(bold=True, size=10), fill=SUB_HEADER_FILL)
    ws.row_dimensions[row_i].height = 22
    row_i += 1

    # 数据行
    data_start = row_i
    for row in time_rows:
        _cell(ws, row_i, 1, _num(row.get("time")))
        for idx, col in enumerate(columns):
            cell = get_cell(row, col.get("id"))
            desc_col = 2 + idx * 2
            _cell(ws, row_i, desc_col, cell.get("desc") or "—")
            ws.cell(row=row_i, column=desc_col + 1).fill = (
                WORK_FILL if cell.get("isWork") else IDLE_FILL
            )
            ws.cell(row=row_i, column=desc_col + 1).alignment = CENTER
        ws.row_dimensions[row_i].height = 28
        row_i += 1

    # 合并连续相同单元格(desc 非空且 desc+isWork 相同)
    for idx, col in enumerate(columns):
        cid = col.get("id")
        desc_col = 2 + idx * 2
        i = 0
        while i < len(time_rows):
            cell = get_cell(time_rows[i], cid)
            span = 1
            if cell.get("desc"):
                while i + span < len(time_rows):
                    nxt = get_cell(time_rows[i + span], cid)
                    if nxt.get("desc") == cell.get("desc") and nxt.get("isWork") == cell.get("isWork"):
                        span += 1
                    else:
                        break
            if span > 1:
                _merge(ws, data_start + i, desc_col, data_start + i + span - 1, desc_col)
                _merge(ws, data_start + i, desc_col + 1, data_start + i + span - 1, desc_col + 1)
                ws.cell(row=data_start + i, column=desc_col).alignment = CENTER
            i += span

    # 统计
    max_time = max((_num(r.get("time")) for r in time_rows), default=0)
    stats = {}
    for col in columns:
        cid = col.get("id")
        work = sum(_num(get_cell(r, cid).get("duration"))
                   for r in time_rows if get_cell(r, cid).get("isWork"))
        idle = max(0, round(max_time - work, 2))
        e = f"{work / max_time * 100:.1f}" if max_time > 0 else "0.0"
        stats[cid] = {"name": col.get("name") or "", "cycleTime": round(max_time, 2),
                      "workTime": round(work, 2), "idleTime": idle, "efficiency": e}

    _merge(ws, row_i, 1, row_i, total_cols)
    _cell(ws, row_i, 1, "统计汇总", font=Font(bold=True, size=12))
    ws.row_dimensions[row_i].height = 26
    row_i += 1

    segs = split_cols(5)

    def write_stat_row(values, *, bold_row=False, colored=False, eff=None):
        nonlocal row_i
        for i in range(5):
            s, e = segs[i]
            _merge(ws, row_i, s, row_i, e)
            cell = _cell(ws, row_i, s, values[i])
            if bold_row:
                cell.font = Font(bold=True)
            if colored:
                if i == 2:
                    cell.font = Font(color=GREEN)
                elif i == 3:
                    cell.font = Font(color=GRAY)
                elif i == 4:
                    try:
                        ev = float(eff)
                    except (TypeError, ValueError):
                        ev = 0
                    if ev >= 80:
                        cell.font = Font(bold=True, color=GREEN)
                    elif ev < 50:
                        cell.font = Font(bold=True, color=RED)
                    else:
                        cell.font = Font(bold=True)
        ws.row_dimensions[row_i].height = 24
        row_i += 1

    for i in range(5):
        s, e = segs[i]
        _merge(ws, row_i, s, row_i, e)
        _cell(ws, row_i, s, ["操作者", "周程/s", "工作时间/s", "空闲时间/s", "效率/%"][i],
              font=bold, fill=HEADER_FILL)
    ws.row_dimensions[row_i].height = 24
    row_i += 1

    for col in columns:
        s = stats[col.get("id")]
        write_stat_row([s["name"], s["cycleTime"], s["workTime"], s["idleTime"],
                        f"{s['efficiency']}%"], colored=True, eff=s["efficiency"])

    if n > 0:
        total_cycle = sum(stats[c.get("id")]["cycleTime"] for c in columns)
        total_work = sum(stats[c.get("id")]["workTime"] for c in columns)
        total_idle = sum(stats[c.get("id")]["idleTime"] for c in columns)
        avg_eff = f"{total_work / total_cycle * 100:.1f}" if total_cycle > 0 else "0.0"
        write_stat_row(["平均", round(total_cycle / n, 2), round(total_work / n, 2),
                        round(total_idle / n, 2), f"{avg_eff}%"], bold_row=True)
    else:
        write_stat_row(["平均", 0, 0, 0, "0.0%"], bold_row=True)

    _border_all(ws, row_i - 1, total_cols)
    name = header.get("taskName") or "未命名"
    return _save(wb), f"联合作业分析_{name}.xlsx"


# ═══════════════════════ 流程程序图 process-flow ═══════════════════════
_PF_TYPES = ["processing", "inspection", "transport", "waiting", "storage"]
_PF_CHARS = {"processing": "〇", "inspection": "□", "transport": "→",
             "waiting": "D", "storage": "▽"}
_PF_LABELS = {"processing": "操作 〇", "inspection": "检查 □", "transport": "搬运 →",
              "waiting": "等待 D", "storage": "储存 ▽"}


def _build_process_flow(data: dict) -> tuple[bytes, str]:
    header = data.get("headerInfo") or {}
    steps = [s for s in (data.get("steps") or []) if isinstance(s, dict)]

    wb, ws = _ws("流程程序图")
    widths = [24, 10, 10, 8, 8, 8, 8, 8, 14]
    _set_widths(ws, widths)
    bold = Font(bold=True)

    _merge(ws, 1, 1, 1, 9)
    _cell(ws, 1, 1, "流程程序图", font=Font(bold=True, size=14))
    ws.row_dimensions[1].height = 30

    # 左侧基本信息 A2:E8
    info_rows = [
        ("工作名称", header.get("workName") or "—"), ("开始", header.get("start") or "—"),
        ("结束", header.get("end") or "—"), ("研究者", header.get("researcher") or "—"),
        ("审阅者", header.get("reviewer") or "—"), ("日期", header.get("date") or "—"),
        ("版本", header.get("version") or "—"),
    ]
    for idx, (k, v) in enumerate(info_rows):
        r = 2 + idx
        _cell(ws, r, 1, k, font=bold, fill=LABEL_FILL_GRAY)
        _merge(ws, r, 2, r, 5)
        _cell(ws, r, 2, v)
        ws.row_dimensions[r].height = 26

    # 右侧统计表 F2:I8
    _merge(ws, 2, 6, 2, 9)
    _cell(ws, 2, 6, "统计表", font=Font(bold=True, size=12), fill=HEADER_FILL)
    for i, h in enumerate(["项别", "次数", "时间/s", "距离/m"]):
        _cell(ws, 3, 6 + i, h, font=bold, fill=HEADER_FILL)

    counts = {t: 0 for t in _PF_TYPES}
    by_type = {t: {"time": 0, "distance": 0} for t in _PF_TYPES}
    for s in steps:
        sym = s.get("symbol")
        if sym in counts:
            counts[sym] += 1
            by_type[sym]["time"] += _num(s.get("time"))
            by_type[sym]["distance"] += _num(s.get("distance"))
    for idx, t in enumerate(_PF_TYPES):
        r = 4 + idx
        for i, v in enumerate([_PF_LABELS[t], counts[t], by_type[t]["time"],
                               by_type[t]["distance"]]):
            _cell(ws, r, 6 + i, v)
        ws.row_dimensions[r].height = 26

    # 工序表头 Row9-10
    _merge(ws, 9, 1, 10, 1)
    _merge(ws, 9, 2, 10, 2)
    _merge(ws, 9, 3, 10, 3)
    _merge(ws, 9, 4, 9, 8)
    _merge(ws, 9, 9, 10, 9)
    _cell(ws, 9, 1, "工作说明")
    _cell(ws, 9, 2, "时间/s")
    _cell(ws, 9, 3, "距离/m")
    _cell(ws, 9, 4, "工序系列")
    _cell(ws, 9, 9, "备注")
    for i, h in enumerate(["操作", "检查", "搬运", "等待", "储存"]):
        _cell(ws, 10, 4 + i, h)
    ws.row_dimensions[9].height = 24
    ws.row_dimensions[10].height = 24

    # 工序数据
    start_row = 11
    for idx, s in enumerate(steps):
        r = start_row + idx
        _cell(ws, r, 1, s.get("content") or "—")
        _cell(ws, r, 2, _num(s.get("time")))
        _cell(ws, r, 3, _num(s.get("distance")))
        sym = s.get("symbol")
        for ci, t in enumerate(_PF_TYPES):
            _cell(ws, r, 4 + ci, _PF_CHARS[t] if sym == t else "")
        _cell(ws, r, 9, s.get("remarks") or "—")
        ws.row_dimensions[r].height = 24

    last_row = start_row + max(len(steps), 1) - 1

    # 工序系列红色连线叠加(D~H 列)
    if len(steps) >= 2:
        sym_widths_px = [_col_px(widths[c - 1]) for c in (4, 5, 6, 7, 8)]
        row_heights_px = [_row_px(24)] * len(steps)
        cols = [_PF_TYPES.index(s.get("symbol")) if s.get("symbol") in _PF_TYPES else 0
                for s in steps]
        png = _polyline_png(sym_widths_px, row_heights_px, cols,
                            color=(239, 68, 68), line_width=2)
        if png:
            _overlay_image(ws, png, tl_col=3, tl_row=start_row - 1,
                           br_col=8, br_row=start_row - 1 + len(steps))

    _border_all(ws, last_row, 9)
    for r in (9, 10):
        for c in range(1, 10):
            ws.cell(row=r, column=c).font = bold
            ws.cell(row=r, column=c).fill = HEADER_FILL

    name = header.get("workName") or "导出"
    return _save(wb), f"流程程序图_{name}.xlsx"


# ═══════════════════════ 动素分析 therblig ═══════════════════════
_THERBLIGS = {
    "TE": ("伸手", 1), "G": ("握取", 1), "TL": ("移物", 1), "P": ("定位", 1),
    "DA": ("拆卸", 1), "U": ("使用", 1), "A": ("装配", 1), "RL": ("放开", 1),
    "I": ("检查", 1),
    "Sh": ("寻找", 2), "F": ("发现", 2), "St": ("选择", 2), "Pn": ("思考", 2),
    "PP": ("预定位", 2),
    "H": ("持住", 3), "AD": ("可避免延迟", 3), "UD": ("不可避免延迟", 3), "R": ("休息", 3),
}
_TB_UPPER = {k.upper(): k for k in _THERBLIGS}
_TB_ALIASES = {"D": "UD"}
# 与 generator.py / TherbligChart.vue 同源（viewBox 0 0 32 32）
_THERBLIG_ICONS = {
    "TE": [("path", "M6 8 Q16 28 26 8")],
    "G": [("path", "M6 26 L6 12 Q6 2 16 2 Q26 2 26 12 L26 26")],
    "TL": [("path", "M4 8 Q16 32 28 8"), ("circle", 16, 17, 3)],
    "P": [("circle", 14, 11, 8), ("path", "M22 11 L22 24 Q22 30 10 30")],
    "DA": [("line", 6, 16, 26, 16), ("line", 13, 9, 11, 23), ("line", 21, 9, 19, 23)],
    "U": [("path", "M8 6 L8 20 Q8 28 16 28 Q24 28 24 20 L24 6")],
    "A": [("line", 6, 12, 26, 12), ("line", 6, 20, 26, 20),
          ("line", 12, 6, 12, 26), ("line", 20, 6, 20, 26)],
    "RL": [("path", "M4 24 Q16 0 28 24"), ("circle", 16, 15, 3)],
    "I": [("ellipse", 16, 16, 8, 11)],
    "Sh": [("path", "M4 18 Q16 6 28 18"), ("path", "M4 18 Q16 30 28 18"),
           ("path", "M22 13 Q17 18 22 23")],
    "F": [("ellipse", 16, 16, 12, 8), ("ellipse", 16, 16, 6, 8)],
    "St": [("line", 4, 16, 26, 16), ("polyline", (20, 10, 26, 16, 20, 22))],
    "Pn": [("circle", 14, 7, 5), ("path", "M9 7 L9 28 L24 28"),
           ("path", "M17 10 L20 13 L11 20 L9 16")],
    "PP": [("circle", 16, 9, 4.5), ("circle", 16, 22, 6.5)],
    "H": [("path", "M6 26 L6 14 Q6 4 16 4 Q26 4 26 14 L26 26"),
          ("line", 2, 26, 30, 26)],
    "AD": [("path", "M8 6 L8 22 L21 22"), ("circle", 24, 22, 3)],
    "UD": [("circle", 26, 16, 3), ("path", "M28 14 L14 4 L4 14 L6 16")],
    "R": [("ellipse", 12, 8, 7, 5), ("path", "M12 13 L12 18 L24 18 L24 28")],
}
_TB_FONT_COLORS = {1: "FF16A34A", 2: "FFD97706", 3: "FFDC2626"}
_TB_ICON_RGB = {1: (34, 197, 94), 2: (245, 158, 11), 3: (239, 68, 68)}
_TB_BG_FILLS = {
    1: PatternFill(fill_type="solid", fgColor="FFDCFCE7"),
    2: PatternFill(fill_type="solid", fgColor="FFFEF3C7"),
    3: PatternFill(fill_type="solid", fgColor="FFFEE2E2"),
}
_TB_ICON_CACHE: dict[str, bytes] = {}
_SVG_TOKEN = re.compile(
    r"[MmLlHhVvCcQqTtSsAaZz]|[-+]?(?:\d*\.\d+|\d+)(?:[eE][-+]?\d+)?"
)


def _tb_lookup(code):
    raw = str(code or "").strip()
    if not raw:
        return None
    up = _TB_ALIASES.get(raw.upper(), raw.upper())
    key = _TB_UPPER.get(up)
    if not key:
        return None
    name, cls = _THERBLIGS[key]
    return key, name, cls


def _q_bez(p0, p1, p2, n=10):
    pts = []
    for i in range(n + 1):
        t = i / n
        u = 1 - t
        pts.append((
            u * u * p0[0] + 2 * u * t * p1[0] + t * t * p2[0],
            u * u * p0[1] + 2 * u * t * p1[1] + t * t * p2[1],
        ))
    return pts


def _svg_path_segments(d: str):
    tokens = _SVG_TOKEN.findall(d)
    i = 0
    cx = cy = 0.0
    last_cmd = ""
    current: list[tuple[float, float]] = []
    segments: list[list[tuple[float, float]]] = []

    def take(n):
        nonlocal i
        out = []
        while len(out) < n and i < len(tokens):
            tok = tokens[i]
            if tok.isalpha():
                break
            out.append(float(tok))
            i += 1
        return out if len(out) == n else None

    while i < len(tokens):
        tok = tokens[i]
        if tok.isalpha():
            last_cmd = tok
            i += 1
        elif not last_cmd:
            i += 1
            continue
        c = last_cmd
        if c in ("M", "m"):
            pair = take(2)
            if not pair:
                break
            x, y = pair
            if c == "m":
                x += cx
                y += cy
            if current:
                segments.append(current)
            current = [(x, y)]
            cx, cy = x, y
            last_cmd = "l" if c == "m" else "L"
        elif c in ("L", "l"):
            pair = take(2)
            if not pair:
                break
            x, y = pair
            if c == "l":
                x += cx
                y += cy
            current.append((x, y))
            cx, cy = x, y
        elif c in ("Q", "q"):
            pts = take(4)
            if not pts:
                break
            x1, y1, x, y = pts
            if c == "q":
                x1 += cx
                y1 += cy
                x += cx
                y += cy
            current.extend(_q_bez((cx, cy), (x1, y1), (x, y))[1:])
            cx, cy = x, y
        else:
            i += 1
    if current:
        segments.append(current)
    return segments


def _therblig_icon_png(code: str) -> bytes | None:
    looked = _tb_lookup(code)
    if not looked:
        return None
    key, _name, cls = looked
    cached = _TB_ICON_CACHE.get(key)
    if cached:
        return cached
    items = _THERBLIG_ICONS.get(key)
    if not items:
        return None
    size, pad = 48, 2
    scale = (size - 2 * pad) / 32.0
    lw = max(2, round(2.4 * scale))
    color = _TB_ICON_RGB[cls]
    img = PILImage.new("RGBA", (size, size), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)

    def xy(px, py):
        return pad + px * scale, pad + py * scale

    for item in items:
        kind = item[0]
        if kind == "path":
            for seg in _svg_path_segments(item[1]):
                if len(seg) >= 2:
                    draw.line([xy(x, y) for x, y in seg], fill=color, width=lw, joint="curve")
        elif kind == "circle":
            _, cx, cy, r = item
            mx, my = xy(cx, cy)
            rr = r * scale
            draw.ellipse((mx - rr, my - rr, mx + rr, my + rr), outline=color, width=lw)
        elif kind == "ellipse":
            _, cx, cy, rx, ry = item
            mx, my = xy(cx, cy)
            draw.ellipse(
                (mx - rx * scale, my - ry * scale, mx + rx * scale, my + ry * scale),
                outline=color, width=lw,
            )
        elif kind == "line":
            _, x1, y1, x2, y2 = item
            draw.line([xy(x1, y1), xy(x2, y2)], fill=color, width=lw)
        elif kind == "polyline":
            pts = [xy(item[1][i], item[1][i + 1]) for i in range(0, len(item[1]), 2)]
            if len(pts) >= 2:
                draw.line(pts, fill=color, width=lw, joint="curve")
    buf = BytesIO()
    img.save(buf, "PNG")
    data = buf.getvalue()
    _TB_ICON_CACHE[key] = data
    return data


def _add_therblig_icon(ws, png: bytes, *, col: int, row: int,
                       col_width: float = 12, row_height: float = 28) -> None:
    """对齐前端 exceljs: tl col+0.15 / row+0.08, 显示 22×22。行列均为 0 基。"""
    img = XLImage(BytesIO(png))
    img.width = 22
    img.height = 22
    anchor = OneCellAnchor()
    anchor._from = AnchorMarker(
        col=col,
        colOff=pixels_to_EMU(0.15 * _col_px(col_width)),
        row=row,
        rowOff=pixels_to_EMU(0.08 * _row_px(row_height)),
    )
    anchor.ext = XDRPositiveSize2D(pixels_to_EMU(22), pixels_to_EMU(22))
    img.anchor = anchor
    ws.add_image(img)


def _build_therblig(data: dict) -> tuple[bytes, str]:
    steps = [s for s in (data.get("steps") or []) if isinstance(s, dict)]

    wb, ws = _ws("动素分析")
    _set_widths(ws, [8, 16, 20, 12, 12, 12, 20, 14])
    bold = Font(bold=True)

    _merge(ws, 1, 1, 1, 8)
    _cell(ws, 1, 1, "动素分析图", font=Font(bold=True, size=14))
    ws.row_dimensions[1].height = 30

    _merge(ws, 2, 1, 3, 1)
    _merge(ws, 2, 2, 3, 2)
    _merge(ws, 2, 3, 3, 3)
    _merge(ws, 2, 4, 2, 6)
    _merge(ws, 2, 7, 3, 7)
    _merge(ws, 2, 8, 3, 8)
    _cell(ws, 2, 1, "序号")
    _cell(ws, 2, 2, "要素作业")
    _cell(ws, 2, 3, "左手动作")
    _cell(ws, 2, 4, "动素记号")
    _cell(ws, 2, 7, "右手动作")
    _cell(ws, 2, 8, "改善点")
    _cell(ws, 3, 4, "左手")
    _cell(ws, 3, 5, "眼")
    _cell(ws, 3, 6, "右手")
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 24
    for r in (2, 3):
        for c in range(1, 9):
            cell = ws.cell(row=r, column=c)
            cell.font = bold
            cell.alignment = CENTER
            cell.fill = HEADER_FILL

    start_row = 4
    for idx, s in enumerate(steps):
        r = start_row + idx
        _cell(ws, r, 1, idx + 1)
        _cell(ws, r, 2, s.get("elementWork") or "")
        _cell(ws, r, 3, s.get("leftAction") or "")
        for col, key in ((4, "leftTherblig"), (5, "eyeTherblig"), (6, "rightTherblig")):
            looked = _tb_lookup(s.get(key))
            if looked:
                code, name, cls = looked
                _cell(ws, r, col, name,
                      font=Font(bold=True, size=8, color=_TB_FONT_COLORS[cls]),
                      fill=_TB_BG_FILLS[cls])
                png = _therblig_icon_png(code)
                if png:
                    _add_therblig_icon(ws, png, col=col - 1, row=r - 1)
        _cell(ws, r, 7, s.get("rightAction") or "")
        _cell(ws, r, 8, s.get("improvement") or "")
        ws.row_dimensions[r].height = 28

    last_data_row = start_row + max(len(steps), 1) - 1

    # 要素作业合并(非空连续相同)
    i = 0
    while i < len(steps):
        span = 1
        if steps[i].get("elementWork"):
            while i + span < len(steps) and \
                    steps[i + span].get("elementWork") == steps[i].get("elementWork"):
                span += 1
        if span > 1:
            _merge(ws, start_row + i, 2, start_row + i + span - 1, 2)
        i += span

    # 统计
    counts = {"left": {1: 0, 2: 0, 3: 0}, "eye": {1: 0, 2: 0, 3: 0},
              "right": {1: 0, 2: 0, 3: 0}}
    for s in steps:
        for side, key in (("left", "leftTherblig"), ("eye", "eyeTherblig"),
                          ("right", "rightTherblig")):
            looked = _tb_lookup(s.get(key))
            if looked:
                counts[side][looked[2]] += 1

    stats_start = last_data_row + 1
    _merge(ws, stats_start, 1, stats_start, 8)
    _cell(ws, stats_start, 1, "动素统计汇总", font=Font(bold=True, size=12))
    ws.row_dimensions[stats_start].height = 26

    hr = stats_start + 1
    _merge(ws, hr, 1, hr, 5)
    _cell(ws, hr, 1, "动素类别")
    _cell(ws, hr, 6, "左手")
    _cell(ws, hr, 7, "眼")
    _cell(ws, hr, 8, "右手")
    for c in range(1, 9):
        cell = ws.cell(row=hr, column=c)
        cell.font = bold
        cell.alignment = CENTER
        cell.fill = HEADER_FILL

    class_rows = [("第一类(有效动素)", "FF22C55E"), ("第二类(辅助动素)", "FFF59E0B"),
                  ("第三类(无效动素)", "FFEF4444")]
    for idx, (name, color) in enumerate(class_rows):
        r = hr + 1 + idx
        _merge(ws, r, 1, r, 5)
        _cell(ws, r, 1, name, font=Font(bold=True, color=color))
        _cell(ws, r, 6, counts["left"][idx + 1])
        _cell(ws, r, 7, counts["eye"][idx + 1])
        _cell(ws, r, 8, counts["right"][idx + 1])
        ws.row_dimensions[r].height = 24

    _border_all(ws, hr + 3, 8)
    return _save(wb), "动素分析图.xlsx"


# ═══════════════════════ 注册表与入口 ═══════════════════════
_BUILDERS = {
    "man-machine": _build_man_machine,
    "two-hand": _build_two_hand,
    "joint-operation": _build_joint_operation,
    "process-flow": _build_process_flow,
    "therblig": _build_therblig,
}

SUPPORTED_TEMPLATES = frozenset(_BUILDERS)


def build_chart_excel(template_id: str, chart_data: dict | None) -> tuple[bytes, str]:
    """按模板 id 生成 Excel。返回 (xlsx 字节, 建议文件名);模板不支持时 ValueError。"""
    builder = _BUILDERS.get(template_id)
    if builder is None:
        raise ValueError(f"模板 {template_id} 不支持 Excel 导出")
    if not isinstance(chart_data, dict):
        chart_data = {}
    return builder(chart_data)
