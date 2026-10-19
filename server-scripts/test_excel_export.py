"""excel_export 服务单测:5 个 IE 模板的服务端 Excel 生成。"""

import io

import pytest
from openpyxl import load_workbook

from app.services.excel_export import SUPPORTED_TEMPLATES, build_chart_excel

CASES = {
    "man-machine": {
        "headerInfo": {"taskName": "钻床钻孔", "worker": "操作工A",
                       "date": "2024-03-20", "analyst": "张工"},
        "machines": [{"id": "m1", "name": "立式钻床"}, {"id": "m2", "name": "攻丝机"}],
        "rows": [
            {"time": 15, "workerDesc": "取工件装夹", "workerState": "work",
             "cells": {"m1": {"desc": "空闲", "state": "idle"},
                       "m2": {"desc": "自动攻丝", "state": "work"}}},
            {"time": 30, "workerDesc": "空闲", "workerState": "idle",
             "cells": {"m1": {"desc": "钻孔", "state": "work"},
                       "m2": {"desc": "空闲", "state": "idle"}}},
        ],
    },
    "two-hand": {
        "headerInfo": {"stationName": "装配工站"},
        "steps": [
            {"leftDesc": "等待", "leftSymbol": "delay",
             "rightSymbol": "transport", "rightDesc": "取件"},
            {"leftDesc": "固定", "leftSymbol": "operation",
             "rightSymbol": "operation", "rightDesc": "装配"},
        ],
    },
    "joint-operation": {
        "headerInfo": {"taskName": "换胎", "operator": "团队", "date": "2024-03-20",
                       "location": "维修站", "remarks": ""},
        "columns": [{"id": "col1", "name": "前顶操作员"}, {"id": "col2", "name": "枪手"},
                    {"id": "col3", "name": "搬运工"}],
        "timeRows": [
            {"time": 0.5, "cells": {
                "col1": {"desc": "顶起", "duration": 0.5, "isWork": True},
                "col2": {"desc": "等待", "duration": 0.5, "isWork": False},
                "col3": {"desc": "准备", "duration": 0.5, "isWork": True}}},
            {"time": 1.0, "cells": {
                "col1": {"desc": "保持", "duration": 0.5, "isWork": True},
                "col2": {"desc": "拧松", "duration": 0.5, "isWork": True},
                "col3": {"desc": "递胎", "duration": 0.5, "isWork": True}}},
        ],
    },
    "process-flow": {
        "headerInfo": {"workName": "零件加工", "start": "08:00", "end": "12:00",
                       "researcher": "张三", "reviewer": "李四",
                       "date": "2024-03-20", "version": "A"},
        "steps": [
            {"content": "原料储存", "distance": 0, "time": 0, "symbol": "storage", "remarks": ""},
            {"content": "搬运", "distance": 20, "time": 300, "symbol": "transport", "remarks": ""},
            {"content": "切割", "distance": 0, "time": 900, "symbol": "processing", "remarks": ""},
        ],
    },
    "therblig": {
        "steps": [
            {"elementWork": "组装", "leftAction": "定位", "leftTherblig": "P",
             "eyeTherblig": "", "rightTherblig": "P", "rightAction": "定位螺栓",
             "improvement": ""},
            {"elementWork": "组装", "leftAction": "等待", "leftTherblig": "H",
             "eyeTherblig": "I", "rightTherblig": "A", "rightAction": "装配",
             "improvement": "改善持住"},
        ],
    },
}


@pytest.mark.parametrize("template_id", sorted(SUPPORTED_TEMPLATES))
def test_build_valid_workbook(template_id):
    content, filename = build_chart_excel(template_id, CASES[template_id])
    assert content[:2] == b"PK"
    assert filename.endswith(".xlsx")
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    # 标题行 + 数据存在
    assert ws.cell(row=1, column=1).value
    assert len(ws.merged_cells.ranges) > 0


def test_overlay_images_present():
    # 双手作业:左右两条连线图;流程程序图:一条红色连线图
    content, _ = build_chart_excel("two-hand", CASES["two-hand"])
    ws = load_workbook(io.BytesIO(content)).active
    assert len(ws._images) == 2
    content, _ = build_chart_excel("process-flow", CASES["process-flow"])
    ws = load_workbook(io.BytesIO(content)).active
    assert len(ws._images) == 1


def test_therblig_embeds_icons():
    content, _ = build_chart_excel("therblig", CASES["therblig"])
    ws = load_workbook(io.BytesIO(content)).active
    assert len(ws._images) == 5
    assert ws.cell(row=4, column=4).value == "定位"
    assert ws.cell(row=5, column=4).value == "持住"


def test_man_machine_stats():
    content, _ = build_chart_excel("man-machine", CASES["man-machine"])
    ws = load_workbook(io.BytesIO(content)).active
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert "统计汇总" in values
    assert "利用率/%" in values


def test_unsupported_template():
    with pytest.raises(ValueError):
        build_chart_excel("gantt", {})


def test_empty_data_does_not_crash():
    for tid in SUPPORTED_TEMPLATES:
        content, _ = build_chart_excel(tid, None)
        assert content[:2] == b"PK"
